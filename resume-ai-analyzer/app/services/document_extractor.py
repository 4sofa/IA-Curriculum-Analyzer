from io import BytesIO

from docx import Document
from pypdf import PdfReader

from app.core.config import settings
from app.core.errors import AppError


def extract_resume_text(file_bytes: bytes, extension: str) -> str:
    validate_signature(file_bytes, extension)

    try:
        if extension == ".pdf":
            return extract_pdf_text(file_bytes)
        if extension == ".docx":
            return extract_docx_text(file_bytes)
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            status_code=422,
            code="text_extraction_failed",
            message="Nao foi possivel ler o conteudo do arquivo enviado.",
            details={"type": exc.__class__.__name__},
        ) from exc

    raise AppError(status_code=415, code="unsupported_file_type", message="Formato de arquivo nao suportado.")


def validate_signature(file_bytes: bytes, extension: str) -> None:
    if extension == ".pdf" and not file_bytes.startswith(b"%PDF"):
        raise AppError(status_code=422, code="invalid_pdf", message="O arquivo nao parece ser um PDF valido.")

    if extension == ".docx" and not file_bytes.startswith(b"PK"):
        raise AppError(status_code=422, code="invalid_docx", message="O arquivo nao parece ser um DOCX valido.")


def extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))

    if reader.is_encrypted:
        raise AppError(status_code=422, code="encrypted_pdf", message="PDFs protegidos por senha nao sao suportados.")

    if len(reader.pages) > settings.max_pdf_pages:
        raise AppError(
            status_code=413,
            code="too_many_pdf_pages",
            message="O PDF possui paginas demais para esta analise.",
            details={"max_pages": settings.max_pdf_pages},
        )

    pages = []
    for index, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(f"Pagina {index}\n{page_text}")

    return "\n\n".join(pages)


def extract_docx_text(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]

    table_lines = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_lines.append(" | ".join(cells))

    return "\n".join(paragraphs + table_lines)
